"""
Date the spreadsheet-era training sessions from Google Sheets edit history.

scripts/fetchSheetRevisions.ts saves every revision of a sheet as an .xlsx.
For each session row we look for the first revision in which any of its
cells already holds the value it has today: that is the day the session was
logged (later typo fixes only move individual cells, and the earliest
untouched cell still pins the date). Rows whose earliest match is the very
first surviving revision predate the kept history (Google prunes old
revisions) and are left undated so the parser interpolates them.

Each new split's sheet started life as a copy of the previous split's
workbook, which is why "row present in the first revision" is not enough:
the copied-in rows held the old split's numbers until they were cleared.

Output is a JSON map the parser consumes:

    { "<ABBR>:<sheet>:<row>": { "date": "YYYY-MM-DD", "revision": "..." } }

Usage:
    python scripts/deriveSessionDates.py <out.json> ABBR=<revisions-folder> [ABBR=<folder> ...]
"""

from __future__ import annotations

import datetime as dt
import json
import sys
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl

LOCAL_TZ = ZoneInfo("America/New_York")


def local_date(modified_time: str) -> dt.date:
    t = dt.datetime.fromisoformat(modified_time.replace("Z", "+00:00"))
    return t.astimezone(LOCAL_TZ).date()


def norm(value: object) -> object:
    if value is None:
        return None
    if isinstance(value, str):
        v = value.strip()
        return v or None
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def snapshot(path: Path) -> dict[str, dict[int, object]]:
    """{ "<sheet>:<row>": {col: value} } for every non-empty data cell."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out: dict[str, dict[int, object]] = {}
    for ws in wb.worksheets:
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        first_col = 2 if header and isinstance(header[0], str) and header[0].strip().lower() == "date" else 1
        for row_idx, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            cells = {c: norm(v) for c, v in enumerate(values, start=1) if c >= first_col and norm(v) is not None}
            if cells:
                out[f"{ws.title}:{row_idx}"] = cells
    wb.close()
    return out


def derive(abbr: str, folder: Path) -> dict[str, dict]:
    manifest = json.loads((folder / "revisions.json").read_text(encoding="utf-8"))
    revisions = sorted(manifest["revisions"], key=lambda r: r["modifiedTime"])
    paths = [(rev, folder / f"{rev['modifiedTime'].replace(':', '-')}__{rev['id']}.xlsx") for rev in revisions]
    paths = [(rev, p) for rev, p in paths if p.exists()]
    if not paths:
        print(f"{abbr}: no revision files in {folder}")
        return {}

    final = snapshot(paths[-1][1])
    matched: dict[str, tuple[str, str]] = {}  # row key -> (modifiedTime, revision id)
    first_rev_keys: set[str] = set()

    for i, (rev, path) in enumerate(paths):
        try:
            snap = snapshot(path)
        except Exception as exc:  # a partial export shouldn't stop the walk
            print(f"  skip {path.name}: {exc}")
            continue
        for key, cells in snap.items():
            if key in matched:
                continue
            final_cells = final.get(key)
            if not final_cells:
                continue
            if any(final_cells.get(c) == v for c, v in cells.items()):
                matched[key] = (rev["modifiedTime"], rev["id"])
                if i == 0:
                    first_rev_keys.add(key)

    dated = {
        f"{abbr}:{key}": {"date": local_date(t).isoformat(), "revision": rid}
        for key, (t, rid) in matched.items()
        if key not in first_rev_keys
    }
    print(
        f"{abbr}: {len(paths)} revisions ({paths[0][0]['modifiedTime'][:10]} .. {paths[-1][0]['modifiedTime'][:10]}), "
        f"{len(final)} rows today, {len(dated)} dated, {len(first_rev_keys)} predate the surviving history"
    )
    return dated


def main() -> None:
    out_path = Path(sys.argv[1])
    result: dict[str, dict] = {}
    for spec in sys.argv[2:]:
        abbr, folder = spec.split("=", 1)
        result.update(derive(abbr, Path(folder)))
    out_path.write_text(json.dumps(result, indent=1, sort_keys=True), encoding="utf-8")
    print(f"wrote {len(result)} dated rows to {out_path}")


if __name__ == "__main__":
    main()
