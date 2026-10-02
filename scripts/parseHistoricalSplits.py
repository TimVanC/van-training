"""
Parse the pre-van.training workout spreadsheets (Gym, God Split, RTK, WNU)
into a flat JSON list of sets that `scripts/importHistoricalLifts.ts` loads
into `historical_lift_sets`.

Workbook conventions (owner's shorthand):
  - Each sheet is one workout day (or one muscle group in the Gym book).
    Row 1 = exercise names, optional row 2 = "sets x rep-range", later rows =
    one training session each.
  - A cell lists the sets for that exercise in the session:
      "85, 85(8), 75(8), 75(7)"  -> 85 at target reps, 85x8, 75x8, 75x7
      "60x4"                      -> 4 sets of 60 at target reps
      "20(15)x3"                  -> 3 sets of 20x15
      "110/99"                    -> drop set (stored, excluded from scoring)
      "6:45 2:25"                 -> leg press plates: 6x45 + 2x25 = 320 lb
      "3 plates"                  -> 3 x 45 per side = 270 lb
      Suffix tokens: S/L (seated/lying) on hamstring curls, S (single-arm)
      elsewhere, DD (dual cable), NH / "no hold". Target reps come from the
      "AxB-C" header (upper bound) or default to 10.
  - Bodyweight lifts (pull-ups, dips): bare numbers are reps at bodyweight,
    "100(8)" is 100 lb of assistance for 8 reps (stored as weight -100), and
    the Gym book writes it the other way round: "6(70)" = 6 reps, 70 assist.
  - Sessions carry no dates except the first few WNU rows, so dates are
    interpolated across each split's known date range and flagged estimated.

Usage: python scripts/parseHistoricalSplits.py <workbook-dir> <out.json>
"""

from __future__ import annotations

import datetime as dt
import json
import re
import statistics
import sys
from pathlib import Path

import openpyxl

SPLITS = [
    {
        "file": "Gym.xlsx",
        "name": "Gym",
        "abbr": "GYM",
        "start": dt.date(2022, 10, 25),
        "end": dt.date(2024, 6, 26),
        "layout": "muscle",  # sheets are muscle groups, col A is a session index
    },
    {
        "file": "God Split.xlsx",
        "name": "God Split",
        "abbr": "GOD",
        "start": dt.date(2024, 9, 24),
        "end": dt.date(2025, 3, 4),
        "layout": "day",
    },
    {
        "file": "RTK (Return of the King) Split.xlsx",
        "name": "RTK (Return of the King)",
        "abbr": "RTK",
        "start": dt.date(2025, 3, 6),
        "end": dt.date(2025, 7, 25),
        "layout": "day",
    },
    {
        "file": "WNU Split (Why not us_).xlsx",
        "name": "WNU (Why not us?)",
        "abbr": "WNU",
        "start": dt.date(2025, 7, 27),
        "end": dt.date(2026, 2, 12),
        "layout": "day",
    },
]

# Rows that are copies of another sheet's numbers (a "reference row" the
# owner pasted when starting a new split), not real sessions.
SKIP_ROWS = {
    ("RTK", "Push", 3),
    ("RTK", "Pull", 3),
    ("RTK", "L&S", 3),
    ("GOD", "Sarms", 3),
    ("GOD", "Legs", 3),
}
SKIP_SHEETS = {("WNU", "Sheet6")}  # normalized duplicate of the per-day sheets
SKIP_COLUMNS = {
    "push ups",
    "jump rope",
    "burn out (cable curl x drag curl)",
    "hip adduction",
    "hip abduction",
    "oblique or sum",
}
BODYWEIGHT_KEYS = ("pull up", "pull-up", "dip", "leg raise", "knee raise")
TRICEPS_SINGLE_ARM_KEYS = ("pushdown", "tricep extension", "overhead")
TRICEPS_SINGLE_ARM_MAX = 35  # below this a pushdown/extension is one-handed

REP_SCHEME_RE = re.compile(r"\s*(\d+)\s*x\s*([\dA-Za-z-]+)\s*$", re.IGNORECASE)


def parse_rep_scheme(text: str) -> int | None:
    """'4x8-10' -> 10, '3x20' -> 20, '3xAMRAP' -> None."""
    m = re.match(r"^\s*\d+\s*x\s*(\d+)(?:\s*-\s*(\d+))?\s*$", text, re.IGNORECASE)
    if not m:
        return None
    return int(m.group(2) or m.group(1))


def clean_header(raw: str) -> tuple[str, int | None]:
    """Split 'Flat DB Bench Press 3x8-10' into (name, target reps)."""
    name = raw.strip()
    target = None
    m = REP_SCHEME_RE.search(name)
    if m:
        target = parse_rep_scheme(m.group(0))
        name = name[: m.start()].strip()
    name = re.sub(r"\s+", " ", name)
    return name, target


def date_cell_to_reps(value: dt.datetime) -> list[int]:
    """Excel turned '10, 9, 7' into 2007-10-09. Recover the rep counts."""
    reps = [value.month, value.day]
    yy = value.year % 100
    if 0 < yy <= 20:
        reps.append(yy)
    return reps


class Parsed:
    def __init__(self) -> None:
        self.rows: list[dict] = []
        self.problems: list[str] = []


def leg_press_weight(token: str) -> float | None:
    """'6:45 2:25' / '6:45s' / '2x25' plate shorthand -> total pounds."""
    total = 0.0
    found = False
    for m in re.finditer(r"(\d+)\s*[:x]\s*(\d+)", token):
        n, w = int(m.group(1)), int(m.group(2))
        if w not in (45, 35, 25, 10, 5):
            return None
        total += n * w
        found = True
    return total if found else None


def tokenize(cell: str, *, is_leg_press: bool) -> list[str]:
    text = cell.strip()
    # Normalise a few hand-written quirks before splitting on commas.
    text = text.replace(" ", " ")
    text = re.sub(r"[^\x00-\x7f]+.*$", "", text)  # stray "�>" garbage
    text = re.sub(r"(\d)\.\s+(\d)", r"\1, \2", text)  # "20. 17x3" -> "20, 17x3"
    text = re.sub(r"\s+\.(\d)", r", \1", text)  # "7, 7 .6" -> "7, 7, 6"
    text = re.sub(r"^\.", "", text)  # ".53x3" -> "53x3"
    text = re.sub(r"(\d)\.(?=,|\s*$)", r"\1", text)  # "6." -> "6"
    text = re.sub(r"\(A\)", "", text)
    # "(8:45 2:25)x2 8:45" and "77x2 77(8)" are missing commas. Leg press
    # cells use spaces inside a plate group ("6:45 2:25"), so leave those.
    text = re.sub(r"(\))\s*x\s*(\d)\s+(?=[\d(])", r"\1x\2, ", text)
    if not is_leg_press:
        text = re.sub(r"([\d)])\s+(?=\d+\s*[(x]|\d+\s*$)", r"\1, ", text)
    # "150(8) L 200 S" -> "150(8) L, 200 S"
    text = re.sub(r"\b([LS])\s+(?=\d)", r"\1, ", text)
    parts = [p.strip() for p in text.split(",")]
    return [p for p in parts if p]


PLATE_DENOMS = {45, 35, 25, 10, 5}


def split_multiplier(token: str, *, is_leg_press: bool) -> tuple[str, int, int | None]:
    """Strip a trailing 'xN': N <= 5 means N sets, larger N means N reps.
    Returns (token, set_count, reps_override)."""
    m = re.search(r"([\d)s])\s*x\s*(\d+)\s*$", token)
    if not m:
        return token, 1, None
    n = int(m.group(2))
    # "6:45 2x25" is a plate group typo for "2:25", not a multiplier.
    if is_leg_press and m.group(1).isdigit() and n in PLATE_DENOMS and not re.search(r"\)\s*x", token):
        return token, 1, None
    rest = token[: m.start(2)].rstrip()
    rest = re.sub(r"\s*x$", "", rest).strip()
    if rest.startswith("(") and rest.endswith(")"):
        rest = rest[1:-1].strip()
    if n <= 5:
        return rest, n, None
    return rest, 1, n


def parse_cell(
    cell: object,
    *,
    exercise: str,
    target_reps: int,
    is_bodyweight: bool,
    is_leg_press: bool,
    problems: list[str],
    where: str,
) -> list[dict]:
    """Return the list of sets in one spreadsheet cell."""
    if cell is None:
        return []
    if isinstance(cell, dt.datetime):
        if not is_bodyweight:
            problems.append(f"{where}: date-typed cell {cell!r} on non-bodyweight lift, skipped")
            return []
        return [
            {"weight": 0.0, "reps": r, "drop": False, "variant": None, "raw": str(cell)}
            for r in date_cell_to_reps(cell)
        ]
    if isinstance(cell, (int, float)):
        if is_bodyweight:
            return [{"weight": 0.0, "reps": int(cell), "drop": False, "variant": None, "raw": str(cell)}]
        if cell <= 0:
            return []
        return [{"weight": float(cell), "reps": target_reps, "drop": False, "variant": None, "raw": str(cell)}]

    text = str(cell)
    low = text.lower()
    if low.strip() in ("n/a", "m", ""):
        return []

    sets: list[dict] = []
    pending_variant_from = 0  # index in `sets` from which a later S/L tag applies
    for token in tokenize(text, is_leg_press=is_leg_press):
        raw_token = token
        variant = None
        note_bits: list[str] = []

        # Trailing descriptor words.
        m = re.search(r"\((lying|seated)\)\s*$", token, re.IGNORECASE)
        if m:
            variant = m.group(1).lower()
            token = token[: m.start()].strip()
        if re.search(r"\bno hold\b", token, re.IGNORECASE):
            note_bits.append("no hold")
            token = re.sub(r"\bno hold\b", "", token, flags=re.IGNORECASE).strip()
        for tag, meaning in (("NH", "no hold"), ("DD", "dual"), ("L", "lying"), ("S", "single")):
            if re.search(rf"\s{tag}\s*$", token):
                token = re.sub(rf"\s{tag}\s*$", "", token).strip()
                if meaning in ("lying",):
                    variant = "lying"
                elif meaning == "single":
                    variant = "seated" if "hamstring" in exercise.lower() else "single"
                elif meaning == "dual":
                    variant = "dual"
                else:
                    note_bits.append(meaning)
        token = token.strip()
        if not token:
            continue

        # Multiplier: "x4" = 4 sets, "x8" (>5) = reps.
        token, count, reps_from_mult = split_multiplier(token, is_leg_press=is_leg_press)
        token = re.sub(r"(?<=[\d)])\s*s$", "", token).strip()  # "320 s", "6:45s"
        if not token:
            continue

        drop = False
        reps = reps_from_mult or target_reps
        weight: float | None = None

        if is_leg_press and re.search(r"\d\s*[:x]\s*\d", token):
            inner = token.strip("() ")
            rm = re.search(r"\((\d+(?:\.\d+)?)\)\s*$", inner)
            if rm:
                reps = int(float(rm.group(1)))
                inner = inner[: rm.start()].strip("() ")
            weight = leg_press_weight(inner)
            if weight is None:
                problems.append(f"{where}: could not read leg press token {raw_token!r}")
                continue
        elif re.match(r"^(\d+)\s*plates?$", token, re.IGNORECASE):
            weight = int(re.match(r"^(\d+)", token).group(1)) * 90.0
        else:
            # Generic "W", "W(R)", "W/W2", "W/W2(R)".
            gm = re.match(
                r"^(-?\d+(?:\.\d+)?)\s*(?:/\s*(\d+(?:\.\d+)?)?)?\s*(?:\((\d+(?:\.\d+)?)?\))?\s*$",
                token,
            )
            if not gm:
                problems.append(f"{where}: unreadable token {raw_token!r} in {text!r}")
                continue
            first = float(gm.group(1))
            second = gm.group(2)
            paren = gm.group(3)
            if "/" in token:
                drop = True
            if is_bodyweight:
                # Pull-ups / dips: bare number = reps, parenthesised number is
                # assistance (or reps when the outer number is the assistance).
                if paren is not None:
                    a, b = first, float(paren)
                    if a <= 20 and b > 20:
                        reps, assist = int(a), b
                    elif a > 20:
                        reps, assist = int(b), a
                    else:
                        reps, assist = int(a), 0.0
                    weight = -assist
                else:
                    if first > 30:
                        problems.append(f"{where}: bodyweight lift with bare {first} ({text!r}), skipped as anomaly")
                        continue
                    weight = 0.0
                    reps = int(first)
            else:
                weight = first
                if paren:
                    reps = int(float(paren))
                elif drop:
                    reps = max(1, target_reps - 2)
                    note_bits.append("drop set, reps estimated")

        if weight is None:
            continue
        if not is_bodyweight and weight <= 0:
            problems.append(f"{where}: non-positive weight in {raw_token!r}, skipped")
            continue
        if reps <= 0:
            problems.append(f"{where}: zero reps in {raw_token!r}, skipped")
            continue

        new_sets = [
            {
                "weight": weight,
                "reps": reps,
                "drop": drop,
                "variant": variant,
                "note": "; ".join(note_bits) or None,
                "raw": raw_token,
            }
            for _ in range(count)
        ]
        if variant:
            # A trailing tag like "... L" covers the untagged sets before it.
            for s in sets[pending_variant_from:]:
                if s["variant"] is None:
                    s["variant"] = variant
            pending_variant_from = len(sets) + len(new_sets)
        sets.extend(new_sets)
    return sets


def finalize_variants(sets: list[dict], exercise: str, carry: dict[str, str | None], abbr: str) -> None:
    """Fill in missing seated/lying (hamstring) and single-arm (triceps) variants.
    The Gym book predates the cable stack the single-arm cutoff is tuned to,
    so only the later splits get the weight-based single-arm guess."""
    low = exercise.lower()
    if abbr == "GYM" and "hamstring" not in low:
        return
    if "hamstring" in low or "leg curl" in low:
        for s in sets:
            if s["variant"] is None:
                if s["weight"] >= 170:
                    s["variant"] = "seated"
                else:
                    s["variant"] = carry.get(exercise) or "seated"
        if sets:
            carry[exercise] = sets[-1]["variant"]
    elif any(k in low for k in TRICEPS_SINGLE_ARM_KEYS) and "single" not in low:
        for s in sets:
            if s["variant"] is None and 0 < s["weight"] < TRICEPS_SINGLE_ARM_MAX:
                s["variant"] = "single"


def interpolate_dates(
    row_indices: list[int],
    known: dict[int, dt.date],
    start: dt.date,
    end: dt.date,
) -> dict[int, tuple[dt.date, bool]]:
    """Give every session row a date, keeping known ones and spreading the rest."""
    out: dict[int, tuple[dt.date, bool]] = {}
    if not row_indices:
        return out
    rows = sorted(row_indices)
    anchors: list[tuple[int, dt.date]] = [(rows[0] - 1, start)] if rows[0] not in known else []
    for r in rows:
        if r in known:
            anchors.append((r, known[r]))
    if rows[-1] not in known:
        anchors.append((rows[-1], end))
    anchors.sort()
    for r in rows:
        if r in known:
            out[r] = (known[r], False)
            continue
        prev = max((a for a in anchors if a[0] < r), key=lambda a: a[0], default=anchors[0])
        nxt = min((a for a in anchors if a[0] > r), key=lambda a: a[0], default=anchors[-1])
        if nxt[0] == prev[0]:
            out[r] = (prev[1], True)
            continue
        frac = (r - prev[0]) / (nxt[0] - prev[0])
        days = (nxt[1] - prev[1]).days * frac
        out[r] = (prev[1] + dt.timedelta(days=round(days)), True)
    return out


def parse_workbook(split: dict, path: Path, parsed: Parsed) -> None:
    wb = openpyxl.load_workbook(path, data_only=True)
    abbr = split["abbr"]
    for ws in wb.worksheets:
        if (abbr, ws.title) in SKIP_SHEETS:
            continue
        header = [c.value for c in ws[1]]
        has_date_col = isinstance(header[0], str) and header[0].strip().lower() == "date"
        first_col = 1 if has_date_col else 0
        scheme_row = [c.value for c in ws[2]] if ws.max_row >= 2 else []
        # Only the undated books (God, RTK) carry a "4x8-10" row under the
        # header; in the dated books row 2 is already data ("90x4" is a set).
        has_scheme_row = not has_date_col and any(
            isinstance(v, str) and parse_rep_scheme(v) is not None for v in scheme_row[first_col:]
        )
        first_data_row = 3 if has_scheme_row else 2

        columns: list[tuple[int, str, int, bool, bool]] = []
        for ci in range(first_col, len(header)):
            raw = header[ci]
            if not isinstance(raw, str) or not raw.strip():
                continue
            name, target = clean_header(raw)
            if name.lower() in SKIP_COLUMNS:
                continue
            if has_scheme_row and isinstance(scheme_row[ci], str):
                target = parse_rep_scheme(scheme_row[ci]) or target
            if ws.title.lower() == "abs" and name.lower() == "cable curls":
                name = "Cable Crunch"  # mislabeled column on the WNU Abs sheet
            low = name.lower()
            columns.append(
                (
                    ci,
                    name,
                    target or 10,
                    any(k in low for k in BODYWEIGHT_KEYS),
                    "leg press" in low,
                )
            )

        # Collect sessions (rows with any set data).
        sessions: dict[int, list[tuple[str, int, bool, bool, object]]] = {}
        known_dates: dict[int, dt.date] = {}
        for r in range(first_data_row, ws.max_row + 1):
            if (abbr, ws.title, r) in SKIP_ROWS:
                continue
            cells = []
            for ci, name, target, bw, lp in columns:
                v = ws.cell(row=r, column=ci + 1).value
                if v is None or (isinstance(v, str) and not v.strip()):
                    continue
                cells.append((name, target, bw, lp, v))
            if not cells:
                continue
            sessions[r] = cells
            if has_date_col:
                dv = ws.cell(row=r, column=1).value
                if isinstance(dv, dt.datetime) and dv.year >= 2020:
                    known_dates[r] = dv.date()

        dates = interpolate_dates(list(sessions), known_dates, split["start"], split["end"])
        carry: dict[str, str | None] = {}
        for r, cells in sorted(sessions.items()):
            date, estimated = dates[r]
            session_index = sorted(sessions).index(r) + 1
            for name, target, bw, lp, value in cells:
                where = f"{abbr}/{ws.title}!r{r} {name}"
                sets = parse_cell(
                    value,
                    exercise=name,
                    target_reps=target,
                    is_bodyweight=bw,
                    is_leg_press=lp,
                    problems=parsed.problems,
                    where=where,
                )
                finalize_variants(sets, name, carry, abbr)
                for i, s in enumerate(sets, start=1):
                    exercise_name = name if not s["variant"] else f"{name} ({s['variant']})"
                    parsed.rows.append(
                        {
                            "split_name": split["name"],
                            "split_abbr": abbr,
                            "day_name": ws.title,
                            "session_key": f"{abbr}:{ws.title}:{r}",
                            "session_index": session_index,
                            "date": date.isoformat(),
                            "date_estimated": estimated,
                            "exercise_name": exercise_name,
                            "set_number": i,
                            "weight": s["weight"],
                            "reps": s["reps"],
                            "target_reps": target,
                            "is_drop_set": s["drop"],
                            "raw_cell": str(value),
                            "notes": s.get("note"),
                            "excluded_reason": None,
                        }
                    )


def flag_outliers(rows: list[dict], problems: list[str]) -> None:
    """A session whose load is 1.5x anything in the surrounding sessions is
    almost always a different machine or a typo (e.g. "80, 100x2" on a cable
    fly that otherwise lives at 43-47). Keep the rows, keep them out of scoring."""
    by_key: dict[tuple[str, str], dict[str, list[dict]]] = {}
    for r in rows:
        if r["weight"] <= 0:
            continue
        by_key.setdefault((r["split_abbr"], r["exercise_name"]), {}).setdefault(r["session_key"], []).append(r)
    for key, sessions in by_key.items():
        if len(sessions) < 3:
            continue
        ordered = sorted(sessions.items(), key=lambda kv: (kv[1][0]["date"], kv[1][0]["session_index"]))
        best = [max(s["weight"] for s in sets) for _, sets in ordered]
        for i, (session_key, sets) in enumerate(ordered):
            neighbours = best[max(0, i - 2) : i] + best[i + 1 : i + 3]
            if not neighbours:
                continue
            ceiling = max(neighbours) * 1.5
            for r in sets:
                if r["weight"] > ceiling:
                    r["excluded_reason"] = f"outlier: {r['weight']:g} vs nearby sessions max {max(neighbours):g}"
                    problems.append(f"{key[0]} {key[1]}: {r['weight']:g} flagged at {session_key} (nearby max {max(neighbours):g})")


def main() -> None:
    src_dir = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("backups/historical-splits")
    out_path = Path(sys.argv[2]) if len(sys.argv) > 2 else src_dir / "historical_lifts.json"
    parsed = Parsed()
    for split in SPLITS:
        parse_workbook(split, src_dir / split["file"], parsed)
    flag_outliers(parsed.rows, parsed.problems)
    out_path.write_text(json.dumps(parsed.rows, indent=1), encoding="utf-8")

    print(f"wrote {len(parsed.rows)} sets to {out_path}")
    by_split: dict[str, set[str]] = {}
    for r in parsed.rows:
        by_split.setdefault(r["split_abbr"], set()).add(r["session_key"])
    for abbr, keys in by_split.items():
        n_sets = sum(1 for r in parsed.rows if r["split_abbr"] == abbr)
        print(f"  {abbr}: {len(keys)} sessions, {n_sets} sets")
    if parsed.problems:
        print(f"\n{len(parsed.problems)} notes:")
        for p in parsed.problems:
            print("  -", p)


if __name__ == "__main__":
    main()
